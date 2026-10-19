"""Excel exports: (1) MOF verification of every downloaded paper, (2) per-reaction performance metrics."""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..extraction.data_checks import flag_text
from ..extraction.reaction_schemas import metric_names
from ..paths import EXCEL_DIR
from .figures import DATA_ROOT, REACTIONS, load
from .science_figures import paper_tags

OUT_ROOT = EXCEL_DIR

IS_MOF = {"verified": "Yes", "needs_review": "Review"}
FILL = {"Yes": "C6EFCE", "Review": "FFEB9C", "No": "FFC7CE", "Not checked": "E7E6E6"}


def _style(ws, df: pd.DataFrame, widths: dict[str, int] | None = None, color_col: str | None = None):
    ws.freeze_panes = "B2" if len(df.columns) > 1 else "A2"
    ws.auto_filter.ref = ws.dimensions
    head_fill = PatternFill("solid", fgColor="1F3864")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for i, col in enumerate(df.columns, start=1):
        vals = df[col].astype(str).head(200)
        q = vals.str.len().quantile(0.9) if len(vals) else 10
        w = min(max(len(str(col)) + 2, (int(q) if pd.notna(q) else 10) + 2), 60)
        if widths and col in widths:
            w = widths[col]
        ws.column_dimensions[get_column_letter(i)].width = max(w, 8)
    if color_col and color_col in df.columns:
        j = list(df.columns).index(color_col) + 1
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=j)
            if cell.value in FILL:
                cell.fill = PatternFill("solid", fgColor=FILL[cell.value])
                cell.font = Font(bold=True)


# --------------------------------------------------------------------------- #
# 1. MOF verification workbook
# --------------------------------------------------------------------------- #
VER_COLS = ["reaction", "doi", "doi_url", "title", "year", "journal", "is_mof_paper", "verification_status",
            "catalyst_class", "mof_role", "mof_identity", "reaction_studied", "matches_expected_reaction",
            "mof_synthesized_or_used_in_experimental", "experimental_electrochemistry_reported", "mof_terms_total",
            "mof_terms_per_10k_chars", "mof_in_title", "mof_terms_abstract", "mof_terms_experimental",
            "cof_terms_total", "reason", "review_note", "synthesis_evidence_quote", "catalyst_evidence_quote",
            "confidence"]


def verification_table(rx: str) -> pd.DataFrame:
    man = pd.read_csv(DATA_ROOT / rx / "manifest.csv")
    vpath = DATA_ROOT / rx / "verification" / f"{rx.lower()}_mof_verification.csv"
    ver = pd.read_csv(vpath) if vpath.exists() else pd.DataFrame(columns=["doi"])
    df = man[["doi", "title", "year", "journal", "status"]].merge(ver.drop(columns=["title", "year", "journal"], errors="ignore"),
                                                                 on="doi", how="left")
    df["reaction"] = rx
    df["doi_url"] = "https://doi.org/" + df["doi"].astype(str)
    vs = df["verification_status"]
    df["is_mof_paper"] = np.where(df["status"] != "ok", "Not checked", vs.map(IS_MOF).fillna("No"))
    df.loc[(df["status"] != "ok"), "verification_status"] = "no_full_text (" + df["status"].astype(str) + ")"
    df.loc[df["status"] == "ok", "verification_status"] = vs
    order = {"Yes": 0, "Review": 1, "No": 2, "Not checked": 3}
    df = df.sort_values(by="is_mof_paper", key=lambda s: s.map(order), kind="stable")
    return df[[c for c in VER_COLS if c in df.columns]]


def export_verification(path: Path | None = None) -> Path:
    path = path or OUT_ROOT / "MOF_verification_all_reactions.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    tables = {rx: verification_table(rx) for rx in REACTIONS if (DATA_ROOT / rx / "manifest.csv").exists()}
    summ = []
    for rx, t in tables.items():
        vc = t["is_mof_paper"].value_counts()
        reasons = t.loc[t["is_mof_paper"] == "No", "verification_status"].value_counts().to_dict()
        summ.append({"reaction": rx, "papers_downloaded": len(t),
                     "Yes (genuine MOF paper)": int(vc.get("Yes", 0)), "Review (borderline)": int(vc.get("Review", 0)),
                     "No (not a MOF paper / wrong reaction)": int(vc.get("No", 0)),
                     "Not checked (no full text)": int(vc.get("Not checked", 0)),
                     "rejection reasons": "; ".join(f"{k}: {v}" for k, v in reasons.items())})
    summary = pd.DataFrame(summ)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        _style(xw.sheets["Summary"], summary)
        allt = pd.concat(tables.values(), ignore_index=True)
        allt.to_excel(xw, sheet_name="All papers", index=False)
        _style(xw.sheets["All papers"], allt, {"title": 50, "reason": 60, "synthesis_evidence_quote": 50,
                                                "catalyst_evidence_quote": 50}, "is_mof_paper")
        for rx, t in tables.items():
            t.to_excel(xw, sheet_name=rx, index=False)
            _style(xw.sheets[rx], t, {"title": 50, "reason": 60, "synthesis_evidence_quote": 50,
                                       "catalyst_evidence_quote": 50}, "is_mof_paper")
    return path


# --------------------------------------------------------------------------- #
# 2. Per-reaction performance workbook
# --------------------------------------------------------------------------- #
def _records(rx: str) -> pd.DataFrame:
    rec, _ = load(rx)
    tags = paper_tags()
    rec["paper_tag"] = rec["doi"].str.lower().map(tags).fillna(rec["paper"])
    rec["doi_url"] = "https://doi.org/" + rec["doi"].astype(str)
    rec["metals"] = rec["node_metals"]
    rec["organic_linker_names"] = rec.get("organic_linkers_raw")
    rec["organic_linker_family"] = rec["organic_linkers"]
    rec["morphology"] = rec["morphology"].where(rec["morphology"].astype(str).str.lower() != "not stated")
    rec["crystallinity"] = rec["crystallinity"].where(rec["crystallinity"].astype(str).str.lower() != "not stated")
    rec["data_flag"] = flag_text(rec)
    if "linker_basis" in rec:
        rec["linker_basis"] = rec["linker_basis"].fillna(rec.get("linker_source"))
    return rec


def _columns(rx: str, rec: pd.DataFrame) -> list[str]:
    metrics = [m for m in metric_names(rx) if m in rec.columns]
    for extra in ("nh3_yield_norm", "nh3_yield_basis"):
        if extra in rec.columns:
            metrics.append(extra)
    cols = (["paper_tag", "reaction", "doi", "doi_url", "title", "year", "journal",
             "catalyst_name", "role", "material_class", "mof_or_precursor_name", "precursor_mof",
             "metals", "mof_node_metals", "n_metals",
             "organic_linker_names", "organic_linker_family", "linker_abbrev", "linker_basis", "linker_check",
             "linker_evidence_quote", "morphology", "crystallinity", "bet_surface_area_m2_g", "synthesis_method",
             "electrolyte", "medium", "reference_electrode", "substrate_or_electrode", "cell_type"]
             + metrics + ["data_flag"] +
            ["stability_duration_h", "stability_cycles", "stability_retention_pct", "stability_condition",
             "stability_note", "source_location", "evidence_quote"])
    return [c for c in cols if c in rec.columns]


def _richest(rec: pd.DataFrame, metrics: list[str]) -> pd.DataFrame:
    d = rec[rec["role"] != "benchmark"].copy()
    d["_n"] = d[metrics].notna().sum(axis=1)
    d = d.sort_values("_n", ascending=False, kind="stable")
    d = d.drop_duplicates(["doi", "catalyst_name", "medium"], keep="first").drop(columns="_n")
    return d.sort_index()


def _group_summary(cat: pd.DataFrame, col: str, metric: str, better: str, label: str) -> pd.DataFrame:
    x = cat.copy()
    x[col] = x[col].fillna("").map(lambda s: [v for v in str(s).split("|") if v])
    x = x.explode(col).dropna(subset=[col])
    x = x[x[col] != ""]
    if x.empty:
        return pd.DataFrame(columns=[label, "n_catalysts", "n_papers"])
    g = x.groupby(col)
    out = pd.DataFrame({"n_catalysts": g.size(), "n_papers": g["doi"].nunique()})
    if metric in x and x[metric].notna().any():
        v = x.dropna(subset=[metric]).groupby(col)[metric]
        out[f"best_{metric}"] = v.min() if better == "low" else v.max()
        out[f"median_{metric}"] = v.median()
        out["n_with_metric"] = v.size()
    out = out.reset_index().rename(columns={col: label})
    return out.sort_values(["n_catalysts"], ascending=False)


HEAD = {"HER": ("overpotential_10_ma_cm2_mv", "low"), "OER": ("overpotential_10_ma_cm2_mv", "low"),
        "ORR": ("half_wave_potential_v", "high"), "N2RR": ("nh3_yield_norm", "high"), "NO3RR": ("nh3_yield_norm", "high")}


def export_reaction(rx: str, path: Path | None = None) -> Path:
    path = path or OUT_ROOT / f"{rx}_performance_metrics.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    rec = _records(rx)
    cols = _columns(rx, rec)
    metrics = [m for m in metric_names(rx) if m in rec.columns]
    cat = _richest(rec, metrics)
    metric, better = HEAD[rx]
    ver = verification_table(rx)
    papers = ver[ver["is_mof_paper"].isin(["Yes", "Review"])][["doi", "doi_url", "title", "year", "journal", "is_mof_paper",
                                                               "catalyst_class", "mof_role", "mof_identity", "reason"]]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        c = cat[cols].sort_values(metric, ascending=(better == "low"), na_position="last") if metric in cat else cat[cols]
        c.to_excel(xw, sheet_name="Catalysts by condition", index=False)
        _style(xw.sheets["Catalysts by condition"], c, {"title": 40, "evidence_quote": 50, "linker_evidence_quote": 50,
                                                         "stability_note": 40, "stability_condition": 30,
                                                         "organic_linker_names": 36, "synthesis_method": 36})
        a = rec[cols]
        a.to_excel(xw, sheet_name="All records", index=False)
        _style(xw.sheets["All records"], a, {"title": 40, "evidence_quote": 50, "linker_evidence_quote": 50,
                                              "stability_note": 40, "organic_linker_names": 36})
        for sheet, col, label in (("Linker summary", "organic_linker_family", "organic_linker"),
                                  ("Metal summary", "metals", "metal"),
                                  ("Morphology summary", "morphology", "morphology")):
            s = _group_summary(cat, col, metric, better, label) if col != "morphology" else \
                _group_summary(cat.assign(morphology=cat["morphology"].astype(str).replace("nan", "")), col, metric, better, label)
            s.to_excel(xw, sheet_name=sheet, index=False)
            _style(xw.sheets[sheet], s)
        papers.to_excel(xw, sheet_name="MOF-verified papers", index=False)
        _style(xw.sheets["MOF-verified papers"], papers, {"title": 55, "reason": 60}, "is_mof_paper")
    return path

